from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = (Path(__file__).resolve().parents[2] / '计划书/食刻有数_商业计划书_排版优化版.docx')
OUT = (Path(__file__).resolve().parents[2] / '计划书/食刻有数_AIC算法创新赛计划书_重构版.docx')
doc = Document(SRC)
body = doc.element.body

def fmt(run, size, font, bold=False):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)

def is_heading(p, level):
    return p.style is not None and p.style.name == f'Heading {level}'

all_paras = list(doc.paragraphs)
ref_para = next(p for p in all_paras if is_heading(p, 1) and p.text.strip() == '参考文献')
chapter_paras = [p for p in all_paras if is_heading(p, 1) and p.text.strip() not in ('摘要', '参考文献') and p.text.strip()]
if len(chapter_paras) != 13:
    raise RuntimeError(f'expected 13 source chapters, found {len(chapter_paras)}')

nodes = list(body)
anchor_indices = [nodes.index(p._p) for p in chapter_paras]
ref_index = nodes.index(ref_para._p)
blocks = {}
for i, start in enumerate(anchor_indices):
    end = anchor_indices[i + 1] if i + 1 < len(anchor_indices) else ref_index
    blocks[i + 1] = nodes[start + 1:end]

new_order = [
    ('第一章 项目概述与赛道定位', [1]),
    ('第二章 场景痛点与需求验证', [3, 2]),
    ('第三章 产品方案与使用流程', [7]),
    ('第四章 AI技术与数据体系', [4, 8]),
    ('第五章 实验、试点与阶段证据', [5]),
    ('第六章 市场与竞争定位', [6]),
    ('第七章 商业模式与运营落地', [9, 10]),
    ('第八章 运营实施与团队分工', [11]),
    ('第九章 财务计划与风险合规', [12, 13]),
    ('第十章 项目计划与结语', []),
]

chapter_nodes = {n for values in blocks.values() for n in values}
chapter_nodes.update(p._p for p in chapter_paras)
for node in chapter_nodes:
    if node.getparent() is body:
        body.remove(node)

insert = list(body).index(ref_para._p)
for title, source_ids in new_order:
    p = doc.add_paragraph(style='Heading 1')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run(title)
    elem = p._p
    body.remove(elem)
    body.insert(insert, elem)
    insert += 1
    for source_id in source_ids:
        for node in blocks[source_id]:
            body.insert(insert, node)
            insert += 1
    if title.startswith('第五章'):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(24)
        p.paragraph_format.line_spacing = 1.5
        p.add_run('本章将经营测算、模型回测与后续授权试点分层说明。现有14个营业日、4类档口、7个菜品和98条记录用于字段贯通、指标复算与方法验证，不替代真实现场试点；正式试点将同步记录备餐、售出、剩余、缺货、毛利、工时和建议采纳情况。')
        elem = p._p
        body.remove(elem)
        body.insert(insert, elem)
        insert += 1
    if title.startswith('第十章'):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(24)
        p.paragraph_format.line_spacing = 1.5
        p.add_run('项目将按“数据接入—小范围试点—指标复核—复制推广”的节奏推进，在每一阶段设置可验收的交付物和退出条件。团队将持续根据真实经营反馈修正模型、流程和服务边界。')
        elem = p._p
        body.remove(elem)
        body.insert(insert, elem)
        insert += 1

for p in list(doc.paragraphs):
    if p.style and p.style.name == 'toc 1':
        p._element.getparent().remove(p._element)
toc = next((p for p in doc.paragraphs if p.text.strip() == '目录'), None)
if toc is not None:
    parent = toc._p.getparent()
    at = list(parent).index(toc._p) + 1
    for title, _ in new_order:
        p = doc.add_paragraph(style='toc 1')
        p.add_run(title + '\t')
        parent.remove(p._p)
        parent.insert(at, p._p)
        at += 1

chapter = 0
sub = 0
for p in doc.paragraphs:
    text = p.text.strip()
    style = p.style.name if p.style else ''
    if style == 'Heading 1' and '章' in text:
        chapter += 1
        sub = 0
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        for run in p.runs:
            fmt(run, 16, 'SimHei', True)
    elif style == 'Heading 2':
        sub += 1
        raw = text
        if raw and raw.split(' ', 1)[0].replace('.', '').isdigit():
            raw = raw.split(' ', 1)[1] if ' ' in raw else raw
        p.text = f'{chapter}.{sub} {raw}'
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.keep_with_next = True
        for run in p.runs:
            fmt(run, 14, 'SimHei', True)
    elif style in ('下标', 'Caption') or text.startswith(('图', '表')):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_together = True
        p.paragraph_format.keep_with_next = True
        for run in p.runs:
            fmt(run, 9, 'SimHei')
    elif text and style not in ('toc 1', 'Heading 1'):
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(6)
        if not p.paragraph_format.first_line_indent:
            p.paragraph_format.first_line_indent = Pt(24)
        for run in p.runs:
            fmt(run, 11, 'SimSun')

for table in doc.tables:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tbl_pr.append(borders)
    for edge, value, size in [('top', 'single', '12'), ('insideH', 'nil', '0'), ('insideV', 'nil', '0'), ('bottom', 'single', '12'), ('start', 'nil', '0'), ('end', 'nil', '0')]:
        element = borders.find(qn('w:' + edge))
        if element is None:
            element = OxmlElement('w:' + edge)
            borders.append(element)
        element.set(qn('w:val'), value)
        element.set(qn('w:sz'), size)
        element.set(qn('w:color'), '000000')
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        if trPr.find(qn('w:cantSplit')) is None:
            trPr.append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                for run in p.runs:
                    fmt(run, 10, 'SimSun')

settings = doc.settings._element
update = settings.find(qn('w:updateFields'))
if update is None:
    update = OxmlElement('w:updateFields')
    settings.append(update)
update.set(qn('w:val'), 'true')
doc.save(OUT)
print(OUT)
