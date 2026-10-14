from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from zipfile import ZipFile, ZIP_DEFLATED
from copy import deepcopy
import math, re, json, csv, hashlib

ROOT=Path(__file__).resolve().parents[2]
SRC=ROOT/'计划书/食刻有数_AIC算法创新赛计划书_终审优化版.docx'
OUT=ROOT/'计划书/食刻有数_AIC算法创新赛计划书_图表精修版.docx'
AS=ROOT/'项目代码/docs/figure_revision_assets'; AS.mkdir(exist_ok=True)
FONT='C:/Windows/Fonts/simhei.ttf'
INK='#202824'; GREEN='#245749'; ORANGE='#ad6035'; PALE='#edf3ef'; GRAY='#66716b'
def font(n): return ImageFont.truetype(FONT,n)
class Canvas:
 def __init__(self,h=660):
  self.im=Image.new('RGB',(1800,h),'white'); self.d=ImageDraw.Draw(self.im)
 def text(self,xy,t,size=40,fill=INK):
  lines=t.split('\n'); x,y=xy
  for k,line in enumerate(lines):
   w=self.d.textlength(line,font=font(size)); self.d.text((x-w/2,y+k*(size+12)),line,font=font(size),fill=fill)
 def box(self,x,y,w,h,t,accent=False,size=40):
  self.d.rounded_rectangle((x,y,x+w,y+h),radius=18,fill=GREEN if accent else PALE,outline=GREEN,width=3)
  lh=len(t.split('\n'))*(size+12)-12
  self.text((x+w/2,y+(h-lh)/2-5),t,size,'white' if accent else INK)
 def arrow(self,a,b,color=GREEN,width=5,dashed=False):
  x,y=a; xx,yy=b
  if dashed:
   L=math.hypot(xx-x,yy-y)
   for k in range(0,int(L),24):
    p=k/L; q=min(k+12,L)/L; self.d.line((x+(xx-x)*p,y+(yy-y)*p,x+(xx-x)*q,y+(yy-y)*q),fill=color,width=width)
  else:self.d.line((x,y,xx,yy),fill=color,width=width)
  ang=math.atan2(yy-y,xx-x); ll=20
  self.d.polygon([(xx,yy),(xx-ll*math.cos(ang-.45),yy-ll*math.sin(ang-.45)),(xx-ll*math.cos(ang+.45),yy-ll*math.sin(ang+.45))],fill=color)
 def save(self,n): self.im.save(AS/n,dpi=(300,300))
def process(n,labels,subs=None,feedback=False):
 c=Canvas(420 if not feedback else 540); count=len(labels); w=(1660-(count-1)*55)/count; y=130
 for i,label in enumerate(labels):
  x=70+i*(w+55); c.box(x,y,w,135,label,accent=i==count-1,size=38 if count<6 else 33)
  if subs and subs[i]:c.text((x+w/2,y+155),subs[i],31,GRAY)
  if i<count-1:c.arrow((x+w+8,y+67),(x+w+47,y+67))
 if feedback:
  c.d.line((1730,285,1730,440,70,440,70,285),fill=GREEN,width=4); c.arrow((70,330),(70,265)); c.text((900,450),'下一周期',31,GRAY)
 c.save(n)

# Each figure has one task. Explanatory prose is added below its caption in Word.
c=Canvas(480)
for i,(a,b) in enumerate([('服务','工具＋经营复盘'),('流程','首批→补餐→停做'),('机制','限时、限量、核销'),('数据','建议→采纳→复盘')]):
 x=70+i*430; c.box(x,85,370,110,a,i==3); c.text((x+185,260),b,32)
c.save('image2.png')
c=Canvas(620)
for i,(a,b) in enumerate([('反浪费与食品安全','状态确认'),('档口成本压力','分批备餐'),('学生优惠需求','限时预约'),('数字台账条件','经营复盘')]):
 y=45+i*145; c.box(80,y,640,100,a,size=38); c.arrow((750,y+50),(1040,y+50)); c.box(1070,y,650,100,b,i==3)
c.save('image3.png')
process('image4.png',['首批备餐','需求校准','分批补餐','停止制作','剩余处置'],['历史记录','实际客流','现场判断','截止时点','安全确认'])
process('image6.png',['公开资料','访谈观察','经营建档','基线记录','策略运行','结果复核'])
process('image11.png',['首期','同城','阶段拓展'],['1所高校／约20档口','10所高校／约200档口','50所高校／约1000档口'])
process('image15.png',['档口提出需求','运营方协调','学校后勤准入','服务启动'],['经营数据确认','场地规则','授权与责任','单档口验证'])
# Five-step cycle rather than a second layer of copy.
c=Canvas(850); nodes=[]
for i,t in enumerate(['经营记录','需求判断','现场执行','限时处置','周期复盘']):
 a=-math.pi/2+i*2*math.pi/5; x=900+510*math.cos(a); y=405+300*math.sin(a);nodes.append((x,y))
 c.d.ellipse((x-135,y-85,x+135,y+85),fill=PALE,outline=GREEN,width=4);c.text((x,y-23),t,42)
for i,a in enumerate(nodes):
 b=nodes[(i+1)%5]; v=(b[0]-a[0],b[1]-a[1]); L=math.hypot(*v);c.arrow((a[0]+v[0]*150/L,a[1]+v[1]*100/L),(b[0]-v[0]*150/L,b[1]-v[1]*100/L))
c.text((900,380),'档口确认',40,GREEN);c.save('image16.png')
c=Canvas(550)
for i,(role,scope) in enumerate([('档口端','经营台账\n建议确认与优惠发布'),('学生端','公开餐品与价格\n预约与核销凭证'),('运营管理端','授权经营汇总\n异常与阶段报告')]):
 x=70+i*580;c.box(x,60,500,100,role,i==0);c.text((x+250,235),scope,37)
c.save('image17.png')
# A single gate makes the exception path explicit, avoiding return-to-sale ambiguity.
c=Canvas(670);c.box(60,215,330,140,'预计剩余');c.arrow((405,285),(560,285))
c.d.polygon([(740,145),(920,285),(740,425),(560,285)],fill=PALE,outline=GREEN,width=4);c.text((740,230),'安全与经营\n条件满足？',38)
c.arrow((930,285),(1080,285));c.text((1000,230),'是',32,GREEN);c.box(1100,215,610,140,'限时发布 → 核销 → 下架',True,35)
c.arrow((740,435),(740,525),ORANGE);c.text((800,460),'否',32,ORANGE);c.box(500,535,480,100,'停止发布并记录',size=37);c.save('image18.png')
c=Canvas(760)
for i,(event,act) in enumerate([('教学日','常规建议'),('考试周','校历修正'),('校园活动','事件标记'),('寒暑假','低频维护'),('极端天气','扩大需求区间')]):
 y=45+i*140;c.box(110,y,600,95,event,size=38);c.arrow((750,y+48),(1050,y+48));c.box(1090,y,600,95,act,i==4,size=38)
c.save('image19.png')
process('image30.png',['基线记录','策略设置','现场执行','结果复核','付费判断'],['5—7天','按品类设定','7—14天','同口径比较','意愿与收益'])
c=Canvas(560)
c.box(660,180,470,140,'经营减损服务',True)
for x,y,t in [(70,80,'档口经营者\n首要付费方'),(70,350,'食堂运营方\n准入与管理'),(1250,215,'学生\n预约与反馈')]:
 c.box(x,y,440,140,t,size=36)
 if x<660:c.arrow((x+450,y+70),(645,250))
 else:c.arrow((x-10,y+70),(1150,250))
c.save('image31.png')
process('image32.png',['经营诊断','基础订阅','专业服务','管理版本'],['免费或低价','99—199元／月','299—599元／月','按档口与范围报价'])
# Provider economics are not customer benefit. Keep the actual numeric example separate.
c=Canvas(620)
for i,(t,num) in enumerate([('年均客单价','2200元'),('可变交付成本','800元'),('贡献毛利','1400元')]):
 x=70+i*580;c.box(x,85,500,100,t,i==2);c.text((x+250,245),num,66,GREEN)
c.text((605,260),'-',60,GRAY);c.text((1185,260),'=',60,GRAY)
c.text((900,440),'贡献毛利率 63.6%',43,GREEN);c.save('image34.png')
process('image35.png',['校内触达','经营诊断','短期试用','订阅服务','续费或退出'])
# Risk coordinates preserved from original; now a single qualitative matrix.
c=Canvas(930);left,top,w,h=270,75,1360,680
for x in range(5):
 for y in range(5):
  fill=['#f2f5f3','#edf1ed','#f4eee4','#f2e2d5','#eed1c4'][min(4,(x+y)//2)]
  c.d.rectangle((left+x*w/5,top+(4-y)*h/5,left+(x+1)*w/5,top+(5-y)*h/5),fill=fill,outline='white',width=3)
for t,x,y in [('食品安全',2,5),('付费验证',4,5),('数据授权',3,4),('交付成本',4,4),('续费',3,3),('预测偏差',3,2)]:
 cx=left+(x-.5)*w/5;cy=top+(5-y+.5)*h/5;c.box(cx-118,cy-38,236,76,t,size=34)
for i in range(5):c.text((left+(i+.5)*w/5,top+h+20),str(i+1),30);c.text((left-50,top+(4-i+.5)*h/5-20),str(i+1),30)
c.text((900,855),'发生可能性（定性等级）',35);c.text((115,375),'影响\n程度',38);c.save('image43.png')
process('image44.png',['识别异常','暂停相关操作','负责人核验','记录与复盘'])
process('image46.png',['经营动作','保留记录','复算指标','周期报告'],['备餐与优惠','数量与成本','剩余率与毛利','授权范围汇总'])
# Model and agent architecture: parallel prediction/rule branches, read-only agent, human action.
c=Canvas(910)
c.box(55,325,315,150,'经营台账\n天气与校历',size=36)
c.box(500,135,400,120,'需求预测模型',size=39);c.box(500,540,400,120,'确定性规则',size=39)
c.arrow((380,395),(480,195));c.arrow((380,395),(480,600))
c.box(1050,325,360,150,'结构化结果',True,size=40)
c.arrow((910,195),(1030,365));c.arrow((910,600),(1030,435))
c.box(1480,135,270,120,'智能体复核',size=34);c.box(1480,540,270,120,'人工确认',True,size=36)
c.arrow((1420,365),(1470,220));c.arrow((1420,435),(1470,560));c.arrow((1615,265),(1615,520));c.text((1630,365),'解释\n提示',31,GRAY)
c.d.line((1615,675,1615,800,215,800,215,490),fill=GREEN,width=4);c.arrow((215,560),(215,490));c.text((900,820),'经营结果回传',33,GRAY);c.save('image53.png')

# Honest financial bars, without promotional headline or repeated interpretation.
c=Canvas(790);left=220;bottom=640;scale=12
for val in range(-5,41,5):
 y=bottom-val*scale;c.d.line((left,y,1700,y),fill='#dae2dd',width=2);c.text((155,y-20),str(val),30)
c.d.line((left,bottom,1700,bottom),fill=GRAY,width=3)
for i,(rev,profit) in enumerate([(1.8,-1.4),(11,.25),(35,9.2)]):
 x=430+i*500
 for off,value,col in [(-90,rev,GREEN),(85,profit,ORANGE)]:
  xx=x+off;yy=bottom-value*scale;c.d.rectangle((xx-65,min(yy,bottom),xx+65,max(yy,bottom)),fill=col);c.text((xx,yy-45 if value>=0 else yy+8),str(value),35,col)
 c.text((x,710),f'第{i+1}年',37)
c.text((310,20),'万元',32,GRAY);c.d.rectangle((1060,30,1090,60),fill=GREEN);c.text((1210,24),'营业收入',32);c.d.rectangle((1400,30,1430,60),fill=ORANGE);c.text((1540,24),'营业利润',32);c.save('image38.png')

# Recompute operational figures from the current CSV, not old raster labels.
records=list(csv.DictReader(open(ROOT/'项目代码/src/main/resources/data/foodwise_operations_14d.csv',encoding='utf-8-sig')))
metrics=[]
for phase in ['基线期','干预期']:
 rr=[r for r in records if r['phase']==phase]
 s={k:sum(float(r[k]) for r in rr) for k in ['prepared_qty','sold_qty','leftover_qty','discount_sold_qty','revenue']}
 s['mape']=sum(abs(float(r['predicted_qty'])-float(r['sold_qty']))/float(r['sold_qty']) for r in rr)/len(rr)*100
 s['leftover_rate']=s['leftover_qty']/s['prepared_qty']*100;s['sellthrough']=s['sold_qty']/s['prepared_qty']*100
 metrics.append(s)
c=Canvas(850);left=190;bottom=640;scale=.16
colors=[GREEN,'#5c9d8c',ORANGE]
for val in range(0,3501,500):
 y=bottom-val*scale;c.d.line((left,y,1710,y),fill='#dbe3de',width=2);c.text((125,y-15),str(val),27)
for i,s in enumerate(metrics):
 for j,k in enumerate(['prepared_qty','sold_qty','leftover_qty']):
  x=480+i*830+(j-1)*185;v=s[k];y=bottom-v*scale;c.d.rectangle((x-63,y,x+63,bottom),fill=colors[j]);c.text((x,y-40),str(int(v)),34,colors[j]);c.text((x,670),['备餐','售出','剩余'][j],31)
 c.text((480+i*830,755),['基线期','策略期'][i],36)
c.text((250,20),'数量（份）',31,GRAY);c.save('image54.png')
c=Canvas(820);left=190;bottom=590;scale=4.9
for val in range(0,101,20):
 y=bottom-val*scale;c.d.line((left,y,1710,y),fill='#dbe3de',width=2);c.text((125,y-16),str(val),29)
for j,k in enumerate(['leftover_rate','sellthrough','mape']):
 center=460+j*510
 for i,s in enumerate(metrics):
  x=center+(i-.5)*150;v=s[k];y=bottom-v*scale;c.d.rectangle((x-55,y,x+55,bottom),fill=GREEN if i==0 else ORANGE);c.text((x,y-40),f'{v:.1f}%',33);c.text((x,620),['基线期','策略期'][i],29)
 c.text((center,715),['剩余率','售罄率','MAPE'][j],36)
c.text((260,25),'百分比（%）',31,GRAY);c.save('image55.png')
days=sorted({r['business_date'] for r in records});daily=[]
for day in days:daily.append({k:sum(float(r[k]) for r in records if r['business_date']==day) for k in ['prepared_qty','sold_qty','leftover_qty']})
c=Canvas(840);left=200;right=1700;bottom=640
for v in range(0,501,100):
 y=bottom-v*1.03;c.d.line((left,y,right,y),fill='#dbe3de',width=2);c.text((130,y-18),str(v),29)
for j,k in enumerate(['prepared_qty','sold_qty','leftover_qty']):
 coords=[]
 for i,s in enumerate(daily):
  x=left+55+i*(right-left-100)/13;y=bottom-s[k]*1.03
  if j==2:c.d.rectangle((x-22,y,x+22,bottom),fill='#d9ac8b')
  else:coords.append((x,y));c.d.ellipse((x-6,y-6,x+6,y+6),fill=colors[j])
 if coords:c.d.line(coords,fill=colors[j],width=5)
for i in range(0,14,2):c.text((left+55+i*(right-left-100)/13,680),days[i][5:],28)
for i,t in enumerate(['备餐量','售出量','剩余量']):c.d.rectangle((500+i*390,40,530+i*390,70),fill=colors[i]);c.text((650+i*390,32),t,32)
c.text((260,35),'数量（份）',31,GRAY);c.text((950,765),'营业日期（2026年7月）',33);c.save('image27.png')

doc=Document(SRC)
original=list(doc.paragraphs)
relations=doc.part.rels
media_para={}
for p in original:
 for b in p._p.xpath('.//*[local-name()="blip"]'):
  rid=b.get(qn('r:embed'));media_para[Path(relations[rid].target_ref).name]=p
def after(p,text,caption=False):
 e=OxmlElement('w:p');p._p.addnext(e);q=Paragraph(e,p._parent);q.style=doc.styles['Normal'];q.text=text
 q.paragraph_format.line_spacing=1.5;q.paragraph_format.space_after=Pt(6);q.paragraph_format.first_line_indent=Pt(22)
 if caption:q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.first_line_indent=Pt(0);q.paragraph_format.keep_with_next=True
 for r in q.runs:
  r.font.name='SimHei' if caption else 'SimSun';r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'SimHei' if caption else 'SimSun');r.font.size=Pt(9 if caption else 11);r.font.color.rgb=RGBColor(0,0,0)
 return q
def remove(p):
 if p._p.getparent() is not None:p._p.getparent().remove(p._p)
removed=['image25.png','image28.png','image29.png','image12.png','image13.png','image14.png','image33.png','image36.png','image45.png','image52.png']
for name in removed:
 p=media_para[name];idx=original.index(p)
 # captions sometimes precede pictures; never remove a non-caption paragraph.
 for q in original[max(0,idx-1):idx+3]:
  if re.match(r'^图\s*\d+-\d+',q.text.strip()):remove(q)
 remove(p)

notes={
'image2.png':'服务、流程、机制和数据分别对应经营复盘、分批制作、安全期内优惠和结果回传。现有原型与台账用于检查方法和流程；创新效果仍需在授权试点中验证，不能把模拟差异直接当成现场成果。',
'image3.png':'政策与校园治理限定安全边界，成本压力推动分批备餐，学生需求影响优惠呈现，数字台账支持周期复盘。这些条件解释服务的切入点，付费意愿仍需单独核验。',
'image4.png':'供需偏差可能在首批备餐和补餐时逐步积累。经营者依据现场客流校准需求，在截止时点停止制作；剩余处置另行核验食品状态，不能以优惠延长安全售卖期。',
'image6.png':'公开资料说明行业背景，访谈和观察识别经营问题，连续台账用于复核动作与结果。三类证据不能相互替代；基线与策略比较须保持菜品、时间窗口和指标口径一致。',
'image11.png':'图中档口数为市场测算假设，与表6-1保持同一范围，并非已服务客户。有效市场还须满足可触达、可记录、存在剩余问题和付费关系明确等条件；不把规划阶段的扩张想象换算为已实现收入。',
'image15.png':'档口、运营方与学校后勤应分别确认经营数据、场地规则和授权边界。权责不清时暂停采集与优惠发布；完成准入后先开展单档口验证，再讨论复制。',
'image16.png':'每天的备餐、售出、剩余和优惠核销进入台账，复盘结果用于下一周期需求判断。档口保留最终确认权，异常客流、设备故障和课程调整需单独记录。',
'image17.png':'学生只查看公开餐品、优惠和本人预约；运营方仅查看授权经营汇总，不接触无关学生明细。档口负责经营数据和动作确认，权限按最小必要字段设置，并保留核销日志。',
'image18.png':'发布前须确认仍处于安全售卖期、餐品状态正常且满足经营条件。任一条件不满足即停止发布并留存记录；只有经负责人重新核验后，才可决定恢复或终止。',
'image19.png':'这些场景不是必然连续发生的时间阶段。教学日采用常规记录，考试周与校园活动加入校历或事件修正，寒暑假降低维护频率，极端天气扩大需求区间并加强人工确认。',
'image30.png':'图示为后续验证安排。先保留5—7天基线记录，再按菜品设置策略并连续观察7—14天；比较剩余、缺货、毛利和操作工时后，由经营者确认付费意愿，不以短期销量变化替代持续效果。',
'image31.png':'档口是首要付费方，运营方提供场景准入和管理协同，学生通过预约、核销和反馈参与。订阅是否持续取决于可核对的客户收益、使用负担和价格匹配，不能仅由平台流量推断。',
'image32.png':'价格均为规划方案，须经真实付款与交付工时校准。基础订阅支持日常记录和建议，专业服务增加参数调整与复盘，管理版本依档口数量和服务范围报价，不要求客户逐级升级。',
'image34.png':'此图核算服务提供方的单位经济性：2200元年均客单价减去800元可变交付成本，得到1400元贡献毛利，贡献毛利率为63.6%。它不同于档口的减损收益；档口收益应另行扣除原价销售挤占、新增操作成本和服务费后判断。',
'image35.png':'试用是评估环节，不等于转化承诺。只有经营者确认收益并愿意付款，才进入订阅；收益不足或使用负担过大时调整范围或退出，不能用续费箭头暗示必然增长。',
'image43.png':'矩阵保留原稿的定性判断等级，用于安排风险处置优先级，并非统计发生概率。食品安全、数据授权和交付成本分别落实到现场核验、最小授权与工时管理；具体措施见表9-3。',
'image44.png':'发现食品状态、数据授权或客流异常时先暂停相关操作，由负责人核验并决定恢复或停止。食品状态不明时不得发布优惠；平台提醒不能替代档口与运营方的现场责任。',
'image46.png':'先保存备餐、售出、剩余、优惠、报损和成本，再按统一口径复算剩余率、售罄率及增量毛利。当前台账支持方法验证，真实减损成效须由授权试点记录与负责人确认共同核验。',
'image53.png':'模型输出需求预测，确定性规则负责边界校验；智能体只读取结构化结果，提供复核、解释和异常提示，不修改数量、金额、风险等级或最终经营动作。档口人工确认后执行，经营结果回传用于复盘；外部服务不可用时保留规则与人工判断。',
'image38.png':'图中为三年经营情景，营业收入依次为1.8、11.0和35.0万元，营业利润为−1.4、0.25和9.2万元。它们是基于付费档口、客单价和成本假设的预测，不是已实现收入或订单承诺。'
}
# Complete caption mapping so missing/pre-image captions are repaired by content identity.
captitles={
'image1.jpeg':'中国高校食堂就餐场景（汕头职业技术学院[9]，非本项目现场）',
'image2.png':'创新要点与经营动作','image3.png':'外部条件与服务切入点','image4.png':'一天经营决策链',
'image20.png':'今日经营总览与趋势','image21.png':'档口与菜品经营管理','image22.png':'备餐决策中心','image23.png':'限时优惠与预约核销','image24.png':'周期经营复盘',
'image6.png':'调研与经营验证证据链','image26.png':'经营数据回传界面','image27.png':'十四日备餐、售出与剩余趋势','image30.png':'后续试点验证安排','image53.png':'预测模型、智能体与人工确认分工',
'image54.png':'经营测算阶段数量对比','image55.png':'阶段指标对比：剩余率、售罄率与预测误差','image11.png':'目标市场分层测算','image15.png':'客户准入与服务启动',
'image31.png':'参与角色与付费关系','image32.png':'服务类型与规划价格','image34.png':'单档口服务的单位经济性','image35.png':'客户评估与续费路径',
'image16.png':'经营动作与数据记录闭环','image17.png':'三端信息权限','image18.png':'限时优惠发布判断','image19.png':'高校经营场景适配',
'image38.png':'三年营业收入与利润情景','image43.png':'主要风险定性矩阵','image44.png':'异常处置与复核','image46.png':'减损结果核验路径'}
chapter=0;cnt={};report=[]
for p in list(doc.paragraphs):
 if p.style.name=='Heading 1':
  if re.match('第[一二三四五六七八九十]+章',p.text):chapter+=1
 if not p._p.xpath('.//*[local-name()="blip"]'):continue
 b=p._p.xpath('.//*[local-name()="blip"]')[0];name=Path(relations[b.get(qn('r:embed'))].target_ref).name
 cnt[chapter]=cnt.get(chapter,0)+1;num=f'{chapter}-{cnt[chapter]}'
 # Remove the old adjacent caption without touching explanations.
 for sibling in [p._p.getprevious(),p._p.getnext()]:
  if sibling is not None and sibling.tag==qn('w:p'):
   if re.match(r'^图\s*\d+-\d+', ''.join(sibling.xpath('.//w:t/text()')).strip()):sibling.getparent().remove(sibling)
 cap=after(p,f'图 {num} {captitles[name]}',True)
 if name in notes:after(cap,notes[name])
 if name in ['image54.png','image55.png']:
  after(cap,'数据来源：项目十四日经营台账，基线期与策略期各49条记录。两阶段差异用于流程与指标验证，尚不能单独归因于模型，也不代表真实档口试点成效。')
 if name=='image27.png':after(cap,'数据来源：项目十四日经营台账，按营业日汇总。趋势展示记录中的数量变化；菜品、客流和优惠变化尚未控制，不能据此认定因果效果。')
 # Match displayed geometry to newly drawn image, with a consistent maximum width.
 if (AS/name).exists():
  im=Image.open(AS/name);ratio=im.height/im.width
  for shape in p._p.xpath('.//*[local-name()="extent"]'):
   shape.set('cx',str(int(Inches(6.0))));shape.set('cy',str(int(Inches(6.0)*ratio)))
  for shape in p._p.xpath('.//*[local-name()="ext"]'):
   if shape.get('cx') is not None:shape.set('cx',str(int(Inches(6.0))));shape.set('cy',str(int(Inches(6.0)*ratio)))
  for crop in p._p.xpath('.//*[local-name()="srcRect"]'):crop.getparent().remove(crop)
 p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(3)
 # Alt text is meaningful provenance, not visible drafting copy.
 for dp in p._p.xpath('.//*[local-name()="docPr"]'):dp.set('descr',captitles[name])
 report.append({'number':num,'media':name,'title':captitles[name],'action':'redrawn' if (AS/name).exists() else 'preserved'})

# Remove stale internal references and redundant figure-introduction narration.
for p in list(doc.paragraphs):
 if p._p.xpath('.//*[local-name()="blip"]'):continue
 t=p.text
 if t.startswith('项目在服务、流程、机制和数据四个维度实现创新'):
  p.text='项目围绕工具与复盘、分批制作、安全期内优惠和结果回传组织服务。四项设计分别对应经营动作，验证重点是建议能否被采用、收益能否被核对。'
 if t.startswith('替代方案比较之外，市场进入还受到'):p.text='除替代方案外，项目还需考虑付费方议价、场景准入和食品安全责任。早期通过小范围验证控制交付成本，再形成可核对案例。'
 if t.startswith('经营者经验、纸笔和Excel是最普遍的替代方案；'):remove(p);continue
 if t.startswith('风险矩阵识别了各类风险的优先级'):remove(p);continue
 if t.startswith('十四日趋势展示了每日备餐'):remove(p);continue
 if t.startswith('5.1 模型回测与异常修正机制'):p.text='5.1 经营测算与指标复核'
 if '依据12.1假设' in p.text:p.text=p.text.replace('依据12.1假设','依据9.1节假设')
 if '见图' in p.text:p.text=re.sub(r'（见图\s*\d+-\d+）','',p.text)
 if re.match(r'^图\s*\d+-\d+',p.text.strip()):
  p.style=doc.styles['Normal'];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True
  for r in p.runs:r.font.name='SimHei';r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'SimHei');r.font.size=Pt(9);r.font.color.rgb=RGBColor(0,0,0)
# Empty heading was included in the old edition.
numeric_map={'2879':'2847','2482':'2473','397份':'374份','238份':'232份','236份':'204份','2484':'2501','159份':'142份','49份优惠':'48份优惠','35559.87':'35283.00','35168.71':'35425.20','13.8%':'13.1%','6.0%':'5.4%','86.2%':'86.9%','94.0%':'94.6%','8.9%':'7.7%','4.5%':'3.3%'}
def reconcile(text):
 for old,new in numeric_map.items():text=text.replace(old,new)
 return text
for p in doc.paragraphs:
 if p._p.xpath('.//*[local-name()="blip"]'):continue
 new=reconcile(p.text)
 if new!=p.text:
  for r in p.runs:r.text=reconcile(r.text)
for table in doc.tables:
 for row in table.rows:
  for cell in row.cells:
   for p in cell.paragraphs:
    for r in p.runs:r.text=reconcile(r.text)
    # Bare quantity cells do not contain the unit.
    if p.text.strip() in {'397','159'}:
     p.text={'397':'374','159':'142'}[p.text.strip()]
     for r in p.runs:r.font.name='SimSun';r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'SimSun');r.font.size=Pt(10)
for p in doc.paragraphs:
 if p.text.startswith('当前记录中，策略期备餐量减少'):
  p.text='按当前台账复算，策略期备餐量减少204份，剩余量减少232份，售出量增加28份，MAPE由7.7%降至3.3%。这些差异仅用于检查指标和流程，尚不能单独归因于模型或证明真实试点成效；仍需在同口径授权记录中控制菜品、客流和操作变化。'
  for r in p.runs:r.font.name='SimSun';r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'SimSun');r.font.size=Pt(11)
for table in doc.tables:
 if table.cell(0,0).text=='指标' and any('变化说明' in c.text for c in table.rows[0].cells):
  for row in table.rows:
   label=row.cells[0].text
   replacement={'售出量':'增加28份','营业收入':'增加142.20元'}.get(label)
   if replacement:
    p=row.cells[-1].paragraphs[0];p.text=replacement
    for r in p.runs:r.font.name='SimSun';r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'SimSun');r.font.size=Pt(10)
for p in list(doc.paragraphs):
 if p.style.name.startswith('Heading') and not p.text.strip() and not p._p.xpath('.//w:sectPr'):remove(p)
# Normalize body pagination while preserving the template cover/abstract geometry.
main_started=False
for p in list(doc.paragraphs):
 if p.style.name=='Heading 1' and p.text.startswith('第一章'):main_started=True
 if not main_started:continue
 if re.match(r'^图\d+-\d+',p.text.strip()):remove(p);continue
 if p._p.xpath('.//w:sectPr'):continue
 if not p.text.strip() and not p._p.xpath('.//*[local-name()="blip"]'):
  remove(p);continue
 if p.style.name=='Heading 1' and re.match('第[一二三四五六七八九十]+章',p.text):
  p.paragraph_format.page_break_before=True;p.paragraph_format.keep_with_next=True
 elif p.style.name.startswith('Heading'):
  p.paragraph_format.keep_with_next=True
 elif p._p.xpath('.//*[local-name()="blip"]') or re.match(r'^图\s+\d+-\d+',p.text):
  p.paragraph_format.keep_with_next=True
 else:
  p.paragraph_format.keep_with_next=False;p.paragraph_format.keep_together=False
  if p.text!='目录':p.paragraph_format.page_break_before=False
doc.save(OUT)
tmp=OUT.with_suffix('.tmp.docx')
with ZipFile(OUT) as zi,ZipFile(tmp,'w',ZIP_DEFLATED) as zo:
 for item in zi.infolist():
  data=zi.read(item.filename);name=Path(item.filename).name
  if item.filename.startswith('word/media/') and (AS/name).exists():data=(AS/name).read_bytes()
  zo.writestr(item,data)
tmp.replace(OUT)
(AS/'figure_changes.json').write_text(json.dumps({'source':SRC.name,'output':OUT.name,'removed':removed,'figures':report},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'output':str(OUT),'figures':len(report),'redrawn':sum(x['action']=='redrawn' for x in report),'removed':len(removed)},ensure_ascii=False))
