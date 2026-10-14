from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re, fitz, json

ROOT=Path(__file__).resolve().parents[2]
SRC=ROOT/'计划书/食刻有数_AIC算法创新赛计划书_图表精修版.docx'
OUT=ROOT/'计划书/食刻有数_AIC算法创新赛计划书_全面核查修订版.docx'
PDF=OUT.with_suffix('.pdf')

def set_run(r,font,size,bold=False):
    r.font.name=font; r.font.size=Pt(size); r.bold=bold; r.font.color.rgb=RGBColor(0,0,0)
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),font)

def set_para_text(p,text,font='SimSun',size=11,bold=False,align=None):
    p.text=text
    for r in p.runs:set_run(r,font,size,bold)
    if align is not None:p.alignment=align

def remove_para(p):
    if p._p.getparent() is not None:p._p.getparent().remove(p._p)

doc=Document(SRC)
original=list(doc.paragraphs)
def rewrite_prefix(prefix,text):
    for p in doc.paragraphs:
        if p.text.startswith(prefix):set_para_text(p,text)
def remove_image(name):
    for p in list(doc.paragraphs):
        bs=p._p.xpath('.//*[local-name()="blip"]')
        if not bs:continue
        if Path(doc.part.rels[bs[0].get(qn('r:embed'))].target_ref).name!=name:continue
        nx=p._p.getnext()
        if nx is not None and re.match(r'^图\s*\d+-\d+', ''.join(nx.xpath('.//w:t/text()'))):nx.getparent().remove(nx)
        remove_para(p)

# These figures repeated tables or a simple formula; preserve the actual content in text/tables.
for name in ['image2.png','image11.png','image32.png','image34.png']:remove_image(name)
rewrite_prefix('本项目以高校食堂档口为研究对象','食刻有数面向高校食堂中小档口，针对首批备餐、补餐、停止制作和安全期内剩余处置提供辅助建议。项目以经营台账、需求判断、人工确认和周期复盘连接日常动作，当前处于学生竞赛规划与原型验证阶段。')
rewrite_prefix('基于14个营业日','项目使用14个营业日、4类档口、7个菜品共98条模拟经营记录进行指标复算和流程验证。基线期剩余率13.1%、售罄率86.9%、MAPE为7.7%；策略期分别为5.4%、94.6%和3.3%。上述差异不是现场经营效果或当前校园模型的独立测试成绩。校园模型使用派生目标，仅供预测管线演示；下一阶段将以授权销量标签和固定对照验证模型、建议采纳及经营收益。商业方案包括订阅、专业服务和运营方管理版本，三年收入情景分别为1.8、11.0和35.0万元。')
rewrite_prefix('需求预测输出区间而不是绝对值','当前原型在点预测周围生成经验范围，该范围尚未进行统计覆盖率校准，不能称为置信区间。备餐和补餐参考还需结合制作周期、历史波动及负责人判断。模型异常时采用规则基线；误差监测窗口与触发阈值将在授权试点中预先固定，现阶段不宣称已验证连续三天触发的经营效果。')
rewrite_prefix('现有记录支持字段定义','现有模拟记录支持字段定义、指标计算、数据导入与页面复盘，其中已有考试周事件标签，但尚不构成真实考试周或极端天气验证。优惠是否挤占原价销售、建议是否持续被采纳，需在授权试点中核验。')
rewrite_prefix('建模流程为预训练','模型验证分三层：模拟台账检验指标与流程；当前校园演示模型检验模型加载与预测管线；后续授权试点检验销量预测和经营收益。不同数据、标签和测试窗口的成绩分别报告，不将台账阶段误差作为机器学习模型效果。真实预测验证须先冻结时间切分、规则基线、指标和缺失处理，再比较模型，并独立校准经验范围。')
rewrite_prefix('系统先检查模型结果','确定性校验比较模型、规则和历史范围；模板摘要整理已计算指标；智能体读取结构化事实并提供解释与异常提示。三者分别承担数值边界、固定摘要和语言复核职责。外部服务不可用时保留规则与人工确认，尚无独立实验可证明智能体提升预测准确率或经营收益。')
rewrite_prefix('差异化主要来自','项目的差异化设计是少字段录入、经营动作提示和周期复盘。品类规则、交付规范和连续授权数据可在后续服务中积累，但当前不视为已经形成的竞争壁垒。')
rewrite_prefix('优势在于校园场景','当前基础是产品原型和可复算流程，限制是长期一手数据、付费证据与现场交付经验不足。市场进入仍受食堂准入、传统系统扩展和食品安全责任约束；早期先验证单档口适用性，再讨论复制。')
rewrite_prefix('试用是评估环节','试用后先核对收益、操作负担与支付意愿。符合条件时订阅或续费；收益不足时调整服务范围；责任不清或无法记录时停止服务。')

# Keep the project-specific model provenance explicit and reader-facing.
for p in list(doc.paragraphs):
    if p.text.startswith('本章将经营测算'):remove_para(p)
    if p.text.startswith('综合政策导向与经营机制'):remove_para(p)
    if p.text.startswith('市场规模和定位判断需要'):remove_para(p)
    if p.text.startswith('模型输出需求预测，确定性规则负责'):remove_para(p)
    if p.text.startswith('此图核算服务提供方'):set_para_text(p,'第二年规划年均客单价为2200元、单档口可变交付成本为850元，贡献毛利为1350元，贡献毛利率为61.4%。这是服务提供方的单位经济性；客户减损收益另行扣除原价销售挤占、新增操作成本和服务费。')
    if p.text.startswith('公开餐饮需求预测研究普遍表明'):set_para_text(p,'历史销量、星期、天气和事件标签是本项目拟比较的需求变量。移动平均与同星期基线作为低复杂度对照，树模型用于探索多变量关系；方法优劣必须在固定数据与测试窗口中验证。')
    if p.text.startswith('从政治环境看') or p.text.startswith('从行业竞争结构看') or p.text.startswith('综合优势、劣势、机会与威胁') or p.text.startswith('上述政策、经济、社会和技术因素'):remove_para(p)

# Combine small subheadings; existing associated paragraphs and tables remain editable.
drop_heads={
'2.2 分时段经营决策与闭餐处置','2.4 食堂运营方管理需求','2.5 学生限时优惠消费偏好','2.7 高校餐饮行业特征分析','2.12 战略环境综合分析',
'4.2 经营建档与数据录入规范','4.3 数据处理链路与结果回传','4.8 经营闭环与责任边界',
'6.4 竞争环境与差异化分析','7.6 客户获取与转化','7.7 客户留存与续费','8.3 备餐决策支持与分批制作','8.7 质量管控指标','8.8 服务标准化与复用'}
renames={'2.1 备餐偏差与剩余成本分析':'2.1 经营决策与供需偏差','2.3 档口经营者核心痛点':'2.2 三类用户需求','2.8 反食品浪费政策与校园治理':'2.3 政策与校园治理','2.9 食堂数字化现状与技术方法参考':'2.4 数字化现状与方法选择','2.10 行业机会与市场进入条件':'2.5 首批适用条件',
'4.1 调研资料来源与适用范围':'4.1 资料来源与经营建档','4.4 核心经营指标与计算口径':'4.2 经营指标与计算口径','4.5 基线方法与需求预测模型':'4.3 预测模型与规则基线','4.6 经营复盘、智能摘要与持续迭代':'4.4 周期复盘与模板摘要','4.7 数据权限与隐私安全':'4.5 权限与现场责任','4.9 建模与智能体联合调度':'4.6 模型、规则与智能体分工','4.10 数据来源与治理':'4.7 数据分层与验证边界',
'6.5 市场定位与竞争壁垒':'6.4 市场定位与进入限制','6.6 客户进入顺序与拓展边界':'6.5 客户进入与拓展边界',
'7.4 单档口收益与收费依据':'7.4 客户收益与服务单位经济性','7.5 推广策略与渠道':'7.5 获客、评估与续费',
'8.4 限时优惠机制与安全确认':'8.3 限时优惠与安全确认','8.5 经营复盘与持续改进机制':'8.4 场景适配与周期复盘','8.6 交付流程与SOP':'8.5 交付流程与质量管理','8.9 团队组织与职责分工':'8.6 团队职责与信息核验'}
for p in list(doc.paragraphs):
    if p.text in drop_heads:remove_para(p)
    elif p.text in renames:set_para_text(p,renames[p.text],'SimHei',14,True)

# Make the active style definitions match the direct formatting used by the document.
for sty_name,font,size in [('Normal','SimSun',11),('Body Text','SimSun',11),('Heading 1','SimHei',16),('Heading 2','SimHei',14)]:
    sty=doc.styles[sty_name];sty.font.name=font;sty.font.size=Pt(size);sty.font.bold=sty_name.startswith('Heading');sty.font.color.rgb=RGBColor(0,0,0)
    sty._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),font)

# Remove duplicated section-level explanations and editorial/instructional phrasing.
texts={
    '项目围绕工具与复盘、分批制作、安全期内优惠和结果回传组织服务。四项设计分别对应经营动作，验证重点是建议能否被采用、收益能否被核对。':'项目围绕经营复盘、分批制作、安全期内优惠和结果回传组织服务，验证重点是建议能否被采用、收益能否被核对。',
    '服务、流程、机制和数据分别对应经营复盘、分批制作、安全期内优惠和结果回传。现有原型与台账用于检查方法和流程；创新效果仍需在授权试点中验证，不能把模拟差异直接当成现场成果。':'现有原型与模拟台账用于检查方法和流程；经营效果仍需在授权试点中验证。',
    '营业结束后，实际售出、剩余、优惠核销和异常情况写入经营台账，系统重新计算核心指标并用于下一周期判断。连续记录稳定后，比较多变量方法是否带来稳定改善；若模型连续失准，则回退到简单基线和人工事件修正。数据反馈形成轻量在线学习机制。 智能经营摘要仅调用后端计算结果和预设模板，帮助经营者识别需求变化、风险菜品和复盘重点；所有经营动作仍由负责人确认。':'营业结束后，实际售出、剩余、优惠核销和异常情况写入台账，系统重新计算核心指标并用于下一周期复盘。智能经营摘要只整理后端计算结果和预设模板，所有经营动作仍由负责人确认。',
    '竞争比较不应只写谁的功能更多，而要看谁真正承担了档口减损这一经营任务。收银系统记录交易，ERP管理采购和库存，优惠平台解决流量与核销；食刻有数聚焦的是把备餐、停止制作、预计剩余处置和复盘放进同一条经营链路。':'收银系统记录交易，ERP管理采购和库存，优惠平台解决流量与核销；食刻有数聚焦备餐、停止制作、预计剩余处置和周期复盘。',
    '除替代方案外，项目还需考虑付费方议价、场景准入和食品安全责任。早期通过小范围验证控制交付成本，再形成可核对案例。':'项目还需考虑付费方议价、场景准入和食品安全责任。早期先通过小范围验证控制交付成本，再形成可核对案例。',
    '图中档口数为市场测算假设，与表6-1保持同一范围，并非已服务客户。有效市场还须满足可触达、可记录、存在剩余问题和付费关系明确等条件；不把规划阶段的扩张想象换算为已实现收入。':'图中档口数为市场测算假设，与表6-1保持同一范围，并非已服务客户；长期背景不计入已实现收入。',
    '图示为后续验证安排。先保留5—7天基线记录，再按菜品设置策略并连续观察7—14天；比较剩余、缺货、毛利和操作工时后，由经营者确认付费意愿，不以短期销量变化替代持续效果。':'验证安排为5—7天基线记录和7—14天策略观察，比较剩余、缺货、毛利和操作工时；短期数量变化不等同于长期经营成效。',
    '本章将经营测算、模型回测与后续授权试点分层说明。现有14个营业日、4类档口、7个菜品和98条记录用于字段贯通、指标复算与方法验证，不替代真实现场试点；正式试点将同步记录备餐、售出、剩余、缺货、毛利、工时和建议采纳情况。':'现有14个营业日、4类档口、7个菜品和98条模拟记录用于字段贯通、指标复算与方法验证，不替代真实现场试点。正式试点将同步记录备餐、售出、剩余、缺货、毛利、工时和建议采纳情况。',
    '系统采用规则判断与机器学习模型双重预测。当前实现以模型结果为主、规则结果为辅；模型未就绪时回退规则判断。模型负责预测售出量，智能体只读取结构化事实，负责独立复核、经营解释和异常提示，不得修改备餐量、金额、风险等级或最终经营动作。':'系统同时保留规则基线与机器学习模型。模型预测售出量，规则负责边界校验；智能体只读取结构化结果，提供复核、解释和异常提示，不修改备餐量、金额、风险等级或最终经营动作。',
}
for p in list(doc.paragraphs):
    if p.text in texts:set_para_text(p,texts[p.text])
    if '经营追踪数据按基线期与策略期划分' in p.text:
        set_para_text(p,p.text.replace('经营追踪数据','模拟经营台账'))
    if '项目十四日经营台账' in p.text:
        set_para_text(p,p.text.replace('项目十四日经营台账','项目十四日模拟经营台账'))
    if p.text.startswith('数据反馈形成轻量在线学习机制'):
        set_para_text(p,p.text.replace('数据反馈形成轻量在线学习机制','周期复盘结果用于后续参数校准'))

# Remove duplicated headings while retaining their figures and readable captions.
for p in list(doc.paragraphs):
    if p.text in ['2.6 需求映射与服务动作匹配','2.11 政策导向与市场机会关联']:remove_para(p)

# Replace calculation-sensitive market value and remove ambiguous overclaiming.
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                if '12.8亿元' in p.text:set_para_text(p,p.text.replace('12.8亿元','1.28亿元'))
                if '长期背景' in p.text and '约5.84万个' in p.text:
                    set_para_text(p,p.text.replace('约5.84万个','约5.84万个').replace('12.8亿元','1.28亿元'))

# Preserve only a responsibility plan; unverified biographies are not evidence.
for table in list(doc.tables):
    if table.cell(0,0).text.strip()=='姓名' and table.cell(0,1).text.strip()=='核心职责':
        while len(table.rows)>4:table._tbl.remove(table.rows[-1]._tr)
        vals=[['职责方向','主要任务','验证材料'],['产品与数据','维护原型、数据口径和模型演示','功能演示、数据字典与复算记录'],['场景调研','核对准入、需求与操作负担','授权记录、访谈与观察材料'],['商业运营','测算成本、组织交付与服务评估','成本台账、交付清单与支付意愿记录']]
        for row,values in zip(table.rows,vals):
            for cell,value in zip(row.cells,values):cell.text=value
    elif table.cell(0,0).text.strip()=='姓名' and table.cell(0,1).text.strip()=='年级':
        table._tbl.getparent().remove(table._tbl)
for p in list(doc.paragraphs):
    if '团队成员简介' in p.text or '团队成员情况' in p.text:remove_para(p)
    if p.text.startswith('团队按产品开发'):
        set_para_text(p,'团队按产品与数据、场景调研、商业运营配置职责。实际成员、年级、专业与经历尚需核验；报名人数和资格以当届官方规则为准。')
    if p.text.startswith('[1] 全国高校商业精英'):set_para_text(p,'参赛方向暂按AI+软件创新组织；最终报名方向、团队人数和提交要求以当届AIC官方通知为准。')
# Keep model provenance consistent with the current model card.
t=doc.tables[6]
vals=[['数据层','当前材料','用途与限制'],['跨域参考资料','历史餐饮需求、天气及食物浪费资料','辅助变量设计；来源、地域与许可逐项核验'],['校园演示模型','2600条大学餐饮记录；原始字段无销量与备餐量','销量、备餐等目标依假设派生，仅演示预测管线'],['模拟经营台账','14个营业日、4类档口、7个菜品、98条记录','复算经营指标与检查流程，不代表真实试点'],['授权现场数据','尚待按档口与周期采集','固定切分与基线，验证预测、采纳与经营收益']]
for row,values in zip(t.rows,vals):
    for cell,value in zip(row.cells,values):cell.text=value
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.text=p.text.replace('滚动回测与运行结果核对','计划开展固定时间切分验证；当前仅管线演示').replace('统一窗口和指标进行方法对照','拟采用统一窗口与指标，尚未完成独立对照').replace('14个营业日、98条记录及统一指标口径','14个营业日、98条模拟记录及统一指标口径')
rewrite_prefix('数据按来源与证据强度分层','资料按参考数据、派生目标演示模型、模拟经营台账与未来授权现场数据分层。当前校园模型原始记录缺少销量和备餐量，目标由假设派生，来源地域、许可及真实性仍待核验；Java与Python特征和推理一致性尚未验收。现阶段不能据此宣称校园销量预测准确率或减损收益。')

# Financial table and prose use one explicit variable-cost assumption.
for p in doc.paragraphs:
    if '价值价格比约为2.3—2.5倍' in p.text:
        set_para_text(p,p.text.replace('800元','850元').replace('合计约500元','合计约500元'))
    if '2200元年均客单价减去800元可变交付成本' in p.text:
        set_para_text(p,p.text.replace('800元','850元').replace('1400元贡献毛利','1350元贡献毛利').replace('63.6%','61.4%'))
# Re-apply fonts and spacing after table/text rewrites.
for p in doc.paragraphs:
    if p.style.name=='Heading 1':
        for r in p.runs:set_run(r,'SimHei',16,True)
        p.paragraph_format.keep_with_next=True;p.paragraph_format.page_break_before=bool(re.match(r'^第[一二三四五六七八九十]+章',p.text))
    elif p.style.name=='Heading 2':
        for r in p.runs:set_run(r,'SimHei',14,True)
        p.paragraph_format.keep_with_next=True
    elif p.text.strip().startswith('图 '):
        for r in p.runs:set_run(r,'SimHei',9,False)
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.paragraph_format.first_line_indent=Pt(0)
    else:
        for r in p.runs:
            if r.text:set_run(r,'SimSun',11,False)
        if p.text.strip() and p.style.name not in ('toc 1','Heading 1','Heading 2'):
            p.paragraph_format.line_spacing=1.5
for table in doc.tables:
    for ri,row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment=1
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.15
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ri==0 else (WD_ALIGN_PARAGRAPH.RIGHT if re.fullmatch(r'[约≈+−\-\d.,%—～万亿元份个所/月年月]+',p.text.strip()) else WD_ALIGN_PARAGRAPH.LEFT)
                for r in p.runs:set_run(r,'SimSun',10,ri==0)

# Renumber visible captions after deletion; repair references to retained captions.
counts={}; mapping={}
for p in doc.paragraphs:
    m=re.match(r'^(图|表)\s*(\d+)-(\d+)\s*(.*)',p.text)
    if m:
        kind,ch,old,title=m.groups();key=(kind,ch);counts[key]=counts.get(key,0)+1
        new=f'{kind}{ch}-{counts[key]}';mapping[f'{kind}{ch}-{old}']=new
        set_para_text(p,new+' '+title,'SimHei',9)
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=Pt(0)
        p.paragraph_format.keep_with_next=kind=='表'
for p in doc.paragraphs:
    if re.match(r'^(图|表)\s*\d+-\d+',p.text):continue
    if re.search(r'[图表]\s*\d+-\d+',p.text):
        text=re.sub(r'([图表])\s*(\d+-\d+)',lambda m:mapping.get(m[1]+m[2],m[0]),p.text)
        if text!=p.text:set_para_text(p,text)
# All tables use only top, header-bottom and bottom borders.
for table in doc.tables:
    for ri,row in enumerate(table.rows):
        trPr=row._tr.get_or_add_trPr()
        if ri==0:
            h=OxmlElement('w:tblHeader');trPr.append(h)
        for cell in row.cells:
            tcPr=cell._tc.get_or_add_tcPr()
            old=tcPr.find(qn('w:tcBorders'))
            if old is not None:tcPr.remove(old)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','bottom','left','right','insideH','insideV']:
                el=OxmlElement('w:'+edge);visible=(edge=='top' and ri==0) or (edge=='bottom' and (ri==0 or ri==len(table.rows)-1))
                el.set(qn('w:val'),'single' if visible else 'nil');el.set(qn('w:sz'),'10' if ri in [0,len(table.rows)-1] else '6');el.set(qn('w:color'),'000000');borders.append(el)
            tcPr.append(borders)
doc.save(OUT)
print(OUT)
