from pathlib import Path
from zipfile import ZipFile
import re,json
import fitz
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_TAB_ALIGNMENT,WD_TAB_LEADER
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
P=ROOT.parent/'计划书/食刻有数_AIC算法创新赛计划书_终审优化版.docx'
PDF=P.with_suffix('.pdf')
d=Document(P);pdf=fitz.open(PDF)
heads=[p.text.strip() for p in d.paragraphs if p.style and p.style.name=='Heading 1' and '章' in p.text]
page_map={}
for title in heads:
    for i,page in enumerate(pdf):
        text=re.sub(r'\s+','',page.get_text())
        if re.sub(r'\s+','',title) in text and '目录' not in text:
            page_map[title]=i+1;break
assert len(page_map)==10,page_map
for p in d.paragraphs:
    if p.style and p.style.name=='toc 1':
        title=p.text.strip()
        p.text=f'{title}\t{page_map[title]}'
        p.paragraph_format.tab_stops.add_tab_stop(Pt(420),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
        for r in p.runs:
            r.font.name='SimSun';r.font.size=Pt(11)
            r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'SimSun')
d.save(P)

ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
from lxml import etree
def media(path):
    with ZipFile(path) as z:
        rel=etree.fromstring(z.read('word/_rels/document.xml.rels'));mp={x.get('Id'):x.get('Target') for x in rel}
        xml=etree.fromstring(z.read('word/document.xml'))
        return {Path(mp[b.get(qn('r:embed'))]).name for b in xml.xpath('.//a:blip',namespaces=ns)}
old=media(ROOT.parent/'计划书/食刻有数_AIC算法创新赛计划书_重构版.docx');new=media(P)
assert old-new=={'image5.png','image37.jpeg'},old-new
assert len(pdf)<=70
assert all(len(p.get_text().strip())>30 for p in pdf)
assert len(heads)==10
assert len(d.inline_shapes)==40
assert len(d.tables)==21
assert d.element.body[-1].tag==qn('w:sectPr')
screen_names={f'image{i}.png' for i in [20,21,22,23,24,26]}
assert screen_names.issubset(new)
with ZipFile(P) as z,ZipFile(ROOT.parent/'计划书/食刻有数_AIC算法创新赛计划书_重构版.docx') as oldzip:
    assert all(z.read('word/media/'+n)==oldzip.read('word/media/'+n) for n in screen_names)
report={
    'pages_before_toc_update':len(pdf),'chapters':10,'tables':21,'body_images':40,
    'removed_scene_photos':sorted(old-new),'preserved_screenshots':sorted(screen_names),
    'redrawn_diagrams':['image44.png','image46.png','image52.png','image53.png'],
    'chapter_pages':page_map,'photo_source':'https://www.stpt.edu.cn/2023/0630/c731a28719/page.htm',
    'photo_boundary':'中国高校官方公开实拍；非本项目现场，无合作或试点证明含义。',
    'ciyuanshen':'已生成食堂候选图，视觉检查后不采用；最终关系图以可控排版重画。',
    'validation_scope':'文件结构、媒体字节、图表排版与渲染检查；未新增或运行模型实验。'
}
(ROOT/'docs/AIC终审验收记录.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
