from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from io import BytesIO
import re
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT.parent/'计划书/食刻有数_AIC算法创新赛计划书_终审优化版.docx'
ASSETS=ROOT/'docs/aic_assets'
ASSETS.mkdir(exist_ok=True)
FONT='C:/Windows/Fonts/simhei.ttf'
green='#245749'; gray='#58635f'; orange='#ac6539'

def diagram(name,title,stages,notes):
    im=Image.new('RGB',(1800,820),'white'); d=ImageDraw.Draw(im)
    f=lambda s:ImageFont.truetype(FONT,s)
    d.text((65,45),title,font=f(52),fill='black')
    d.line((65,127,1735,127),fill='#d6dcd9',width=2)
    n=len(stages); gap=38; w=(1670-(n-1)*gap)//n
    for i,(label,lines) in enumerate(stages):
        x=65+i*(w+gap)
        d.line((x,226,x+w,226),fill=green,width=7)
        d.text((x,163),f'{i+1:02d}',font=f(38),fill=green)
        d.text((x,265),label,font=f(37),fill='black')
        for j,line in enumerate(lines):d.text((x,345+j*58),line,font=f(29),fill=gray)
        if i<n-1:
            y=285; xx=x+w+10
            d.line((xx,y,xx+22,y),fill=green,width=3)
            d.polygon([(xx+22,y),(xx+13,y-7),(xx+13,y+7)],fill=green)
    d.line((65,650,1735,650),fill='#d6dcd9',width=2)
    for j,line in enumerate(notes):d.text((65,680+j*48),line,font=f(28),fill=orange if j==0 else gray)
    im.save(ASSETS/name,dpi=(300,300))

diagram('image44.png','异常处置与复核',[
    ('识别异常',['食品状态','数据授权或客流']),('暂停相关操作',['停止优惠发布','隔离异常记录']),('负责人确认',['核验餐品与权限','决定恢复或停止']),('记录并复盘',['保存处理经过','调整下一周期规则'])],
    ['食品安全状态不明时，不进入优惠售卖流程。','平台提供提醒；档口与运营方承担现场确认责任。'])
diagram('image46.png','从经营动作到可核对结果',[
    ('执行动作',['分批备餐与停做','限时优惠与核销']),('保留记录',['备餐、售出、剩余','优惠、报损、成本']),('复算指标',['剩余率与售罄率','避免成本与增量毛利']),('形成复盘',['单档口经营报告','授权范围内汇总'])],
    ['复盘结论用于调整下一周期备餐与处置。','当前台账仅支持方法验证；真实成效需以授权试点记录核验。'])
diagram('image52.png','经营数据如何进入现场决策',[
    ('整理记录',['历史销量与事件','成本和食品状态']),('判断需求',['模型预测区间','规则与异常校验']),('确认执行',['负责人确认数量','备餐、补餐与停做']),('回传结果',['实际售出与剩余','记录异常和采纳'])],
    ['数据 → 判断 → 人工确认 → 经营记录','食品安全判断与最终经营决策保留人工责任。'])
diagram('image53.png','预测模型与智能体的分工',[
    ('数据输入',['经营台账','天气、星期与事件']),('预测与校验',['需求预测模型','确定性规则复核']),('解释与提示',['智能体读取结构化结果','生成原因和异常提示']),('人工决策',['档口确认经营动作','结果进入周期复盘'])],
    ['智能体不修改备餐量、金额、风险等级或最终经营动作。','外部服务不可用时，保留规则校验与人工确认。'])

doc=Document(OUT)
# Previous node moves left the final section-properties element in the middle.
for sp in list(doc.element.body.findall(qn('w:sectPr'))):
    doc.element.body.remove(sp)
    doc.element.body.append(sp)
# Make template text boxes consistent with the AIC-facing edition.
for t in doc.element.xpath('.//w:t'):
    if t.text=='创业计划赛道':t.text='算法创新赛 · AI+场景创新（拟选）'
    elif t.text=='参赛组别：':t.text='项目阶段：'
    elif t.text=='创客组':t.text='规划与原型验证'
def setfont(r,font,size,bold=False):
    r.font.name=font;r.font.size=Pt(size);r.bold=bold;r.font.color.rgb=RGBColor(0,0,0)
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),font)
def put(p,text):
    p.text=text

# Correct repetitions and implementation narration without changing evidence numbers.
for p in doc.paragraphs:
    t=p.text.strip()
    if t.startswith('经营记录经表单整理后录入经营数据表'):
        put(p,'平台将经营台账、需求预测、备餐建议、限时优惠和周期复盘连接起来。档口录入每日经营记录后，可查看建议并确认具体动作；运营方在授权范围内查看汇总与异常。现阶段优先保障少量档口的完整使用流程，再根据实际使用量调整服务容量。')
    elif t.startswith('围绕备餐、剩余、限时处置和周期复盘构成'):
        put(p,'服务覆盖备餐、剩余处置和周期复盘。档口围绕少量核心字段完成记录，平台与服务人员将记录整理为当天可执行的建议，并在周期结束后核对经营结果。')
    elif t.startswith('学生端采用限时限量展示'):
        put(p,'学生端同时展示原价、优惠价、可预约数量、取餐地点和截止时间。预约成功后生成核销凭证，到期自动下架。经营者先确认餐品状态，再发布优惠；价格优惠不改变食品安全要求。')
    elif t.startswith('现有材料中的数据集名称及样本量存在版本差异'):
        put(p,'阶段结论应对应同一数据版本和计算口径。数据字典、模型说明与指标计算记录共同用于复核，避免将不同版本的指标直接比较。')
    elif t.startswith('从当前记录看，策略期在备餐量减少'):
        put(p,'当前记录中，策略期备餐量减少236份、剩余量减少238份，售出量基本稳定，MAPE由8.9%降至4.5%。这些差异可用于检查计算口径和经营流程，尚不能单独归因于模型或证明真实试点成效；仍需控制菜品、客流和记录方式后核验。')
    if '经营数据数据' in p.text and not p._p.xpath('.//a:blip'):
        put(p,p.text.replace('经营数据数据','经营资料'))
    if t.startswith('图1-1'):
        put(p,'图1-1 中国高校食堂就餐场景（引自汕头职业技术学院[9]，非本项目现场）')
    if re.match(r'^图\d+-\d+\s+\d+\s+',t):
        put(p,re.sub(r'^(图\d+-\d+)\s+\d+\s+',r'\1 ',p.text))
    if t.startswith('参赛赛道：'):
        put(p,'拟选方向：算法创新赛 · AI+场景创新')
    if t.startswith('参赛组别：'):
        put(p,'项目阶段：学生竞赛规划与原型验证')

# Repair broken sentence in the source.
ps=list(doc.paragraphs)
for a,b in zip(ps,ps[1:]):
    if a.text.endswith('分别核') and b.text.strip()=='对。':
        put(a,a.text+b.text); b._p.getparent().remove(b._p)

toc=next(p for p in doc.paragraphs if p.text.strip()=='目录')
toc.paragraph_format.page_break_before=True
toc.paragraph_format.keep_with_next=True
for p in doc.paragraphs:
    if p.style and p.style.name=='toc 1':
        p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(8)
        p.paragraph_format.line_spacing=1.15;p.paragraph_format.keep_with_next=False
        p.paragraph_format.first_line_indent=Pt(0)
        for r in p.runs:setfont(r,'SimSun',11)

# Fix mixed image/text paragraphs and keep pictures with their captions.
from docx.text.paragraph import Paragraph
for p in list(doc.paragraphs):
    if not p._p.xpath('.//a:blip'):continue
    if p.text.strip():
        new=OxmlElement('w:p');p._p.addprevious(new);np=Paragraph(new,p._parent)
        np.add_run(p.text)
        for r in list(p._p):
            if r.tag==qn('w:r') and not r.xpath('.//a:blip'):p._p.remove(r)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Pt(0)
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.keep_together=True
    p.paragraph_format.space_after=Pt(3)
    for shape in p._p.xpath('.//wp:extent'):
        cx,cy=int(shape.get('cx')),int(shape.get('cy'))
        if cy>Inches(3.8):
            ratio=Inches(3.8)/cy
            shape.set('cx',str(int(cx*ratio)));shape.set('cy',str(int(cy*ratio)))
            for ext in p._p.xpath('.//a:xfrm/a:ext'):
                ext.set('cx',str(int(cx*ratio)));ext.set('cy',str(int(cy*ratio)))

for p in doc.paragraphs:
    s=p.style.name if p.style else '';t=p.text.strip()
    if s=='Heading 1':
        for r in p.runs:setfont(r,'SimHei',16,True)
        p.paragraph_format.keep_with_next=True
    elif s=='Heading 2':
        for r in p.runs:setfont(r,'SimHei',14,True)
        p.paragraph_format.keep_with_next=True
    elif re.match(r'^[图表]\s*\d+-\d+',t):
        for r in p.runs:setfont(r,'SimHei',9)
        p.paragraph_format.keep_with_next=t.startswith('表')
        p.paragraph_format.first_line_indent=Pt(0)
        p.paragraph_format.keep_together=True
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    elif t and s!='toc 1':
        for r in p.runs:setfont(r,'SimSun',11)
        p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent=Pt(22)
        p.paragraph_format.line_spacing=1.5

for table in doc.tables:
    bp=table._tbl.tblPr
    old=bp.find(qn('w:tblBorders'))
    if old is not None:bp.remove(old)
    border=OxmlElement('w:tblBorders');bp.append(border)
    for edge in ['top','bottom','left','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single' if edge in ('top','bottom') else 'nil');e.set(qn('w:sz'),'12');border.append(e)
    for ri,row in enumerate(table.rows):
        tr=row._tr.get_or_add_trPr()
        if tr.find(qn('w:cantSplit')) is None:tr.append(OxmlElement('w:cantSplit'))
        if ri==0 and tr.find(qn('w:tblHeader')) is None:tr.append(OxmlElement('w:tblHeader'))
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr();old=cp.find(qn('w:tcBorders'))
            if old is not None:cp.remove(old)
            cb=OxmlElement('w:tcBorders');cp.append(cb)
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single' if ri==0 and edge=='bottom' else 'nil');e.set(qn('w:sz'),'6');cb.append(e)
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(3)
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ri==0 else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:setfont(r,'SimSun',10,ri==0)

# Also normalize nested/text-box tables excluded by Document.tables.
for tbl in doc.element.xpath('.//w:tbl'):
    pr=tbl.find(qn('w:tblPr'))
    if pr is None:pr=OxmlElement('w:tblPr');tbl.insert(0,pr)
    for tag in ['tblStyle','tblBorders','shd']:
        for e in list(pr.findall(qn('w:'+tag))):pr.remove(e)
    bo=OxmlElement('w:tblBorders');pr.append(bo)
    for edge in ['top','bottom','left','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single' if edge in ('top','bottom') else 'nil');e.set(qn('w:sz'),'12');bo.append(e)
    rows=tbl.findall(qn('w:tr'))
    for ri,row in enumerate(rows):
        for cell in row.findall(qn('w:tc')):
            cp=cell.find(qn('w:tcPr'))
            if cp is None:cp=OxmlElement('w:tcPr');cell.insert(0,cp)
            for tag in ['tcBorders','shd']:
                for e in list(cp.findall(qn('w:'+tag))):cp.remove(e)
            cb=OxmlElement('w:tcBorders');cp.append(cb)
            for edge in ['top','bottom','left','right','insideH','insideV']:
                v='single' if (edge=='top' and ri==0) or (edge=='bottom' and ri in (0,len(rows)-1)) else 'nil'
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),v);e.set(qn('w:sz'),'6' if ri==0 and edge=='bottom' else '12');cb.append(e)

# Clear fragile WPS text-box footers and use native Word page fields.
for section in doc.sections:
    for footer in [section.footer,section.first_page_footer,section.even_page_footer]:
        for child in list(footer._element):footer._element.remove(child)
        p=footer.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for value in ['第 ', 'PAGE',' 页']:
            r=p.add_run()
            if value=='PAGE':
                f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),' PAGE ');r._r.addnext(f)
            else:r.text=value
            setfont(r,'SimSun',9)
    pn=section._sectPr.find(qn('w:pgNumType'))
    if pn is not None:section._sectPr.remove(pn)

ref=doc.add_paragraph('[9] 汕头职业技术学院. 暖心陪餐：吴萍校长带队深入学生食堂，与学生共享美好“食”光[EB/OL]. 2023-06-19. https://www.stpt.edu.cn/2023/0630/c731a28719/page.htm（访问日期：2026-10-10）. 图片来源：总务处。仅作场景引用，不代表项目合作或试点。')
for r in ref.runs:setfont(r,'SimSun',10)
doc.save(OUT)
tmp=OUT.with_suffix('.final.tmp.docx')
with ZipFile(OUT) as zin,ZipFile(tmp,'w',ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data=zin.read(item.filename)
        name=Path(item.filename).name
        if item.filename=='word/media/image1.jpeg':data=(ASSETS/'chinese_campus.jpg').read_bytes()
        elif item.filename.startswith('word/media/') and name in ['image44.png','image46.png','image52.png','image53.png']:data=(ASSETS/name).read_bytes()
        zout.writestr(item,data)
tmp.replace(OUT)
print('Saved',OUT)
